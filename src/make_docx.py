# -*- coding: utf-8 -*-
"""draft_full_v2.md → paper/draft_review_v1.docx (공동연구자 검토용).
줄번호·더블스페이싱·넓은 여백·검토메모·표·Figure placeholder·참고문헌 분리.
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_LINE_SPACING, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
SRC_MD = ROOT / "paper" / "draft_full_v2.md"
OUT = ROOT / "paper" / "draft_review_v1.docx"

doc = Document()

# ---------- 기본 스타일: 폰트(한글 Malgun / 영문 Times), 더블 스페이싱 ----------
def style_font(name, ascii_="Times New Roman", ea="Malgun Gothic", size=11):
    try:
        st = doc.styles[name]
    except KeyError:
        return
    st.font.size = Pt(size)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:ascii"), ascii_); rf.set(qn("w:hAnsi"), ascii_)
    rf.set(qn("w:eastAsia"), ea)

for s in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
    style_font(s)
normal = doc.styles["Normal"]
normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE

# ---------- 여백 넓게 + 줄번호 ----------
sec = doc.sections[0]
sec.top_margin = Inches(1.0); sec.bottom_margin = Inches(1.0)
sec.left_margin = Inches(1.25); sec.right_margin = Inches(2.2)   # 오른쪽 넓게(코멘트용)
sectPr = sec._sectPr
ln = OxmlElement("w:lnNumType")
ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); ln.set(qn("w:distance"), "360")
pgMar = sectPr.find(qn("w:pgMar"))
pgMar.addnext(ln)

# ---------- 인라인 서식(**bold** *italic* `code`) ----------
def add_runs(p, text):
    for tok in re.split(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            p.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = p.add_run(tok[1:-1]); r.font.name = "Consolas"
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            p.add_run(tok[1:-1]).italic = True
        else:
            p.add_run(tok)

def clean(text):
    return text.replace("**", "").replace("`", "").strip()

# ---------- md 읽기 ----------
lines = SRC_MD.read_text(encoding="utf-8").split("\n")

# 프런트매터: 제목 + 메타
title, meta, i = None, [], 0
while i < len(lines):
    if lines[i].startswith("# ") and not lines[i].startswith("## "):
        title = lines[i][2:].strip(); i += 1; break
    i += 1
while i < len(lines):
    s = lines[i].strip()
    if s.startswith("## "):
        break
    if s.startswith("**") and ":**" in s:
        meta.append(s)
    i += 1

# ---------- 표지 ----------
t = doc.add_heading(title, level=0)
t.alignment = WD_ALIGN_PARAGRAPH.CENTER

au = doc.add_paragraph(); au.alignment = WD_ALIGN_PARAGRAPH.CENTER
au.add_run("Author 1¹, Author 2²").bold = True
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.add_run("(공동연구자 성명·소속 삽입 위치 — placeholder)").italic = True

for m in meta:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    add_runs(p, m)

# 검토 요청 메모 박스 (음영 1x1 표)
doc.add_paragraph()
box = doc.add_table(rows=1, cols=1); box.style = "Table Grid"
cell = box.rows[0].cells[0]
shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "FFF2CC")
cell._tc.get_or_add_tcPr().append(shd)
h = cell.paragraphs[0]; h.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
h.add_run("■ 검토 요청 사항").bold = True
for item in ["(1) 연구 격차 및 가설의 타당성",
             "(2) 결과 해석의 적절성",
             "(3) 한계 섹션 보완 필요 여부"]:
    q = cell.add_paragraph(item)
    q.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    q.paragraph_format.left_indent = Inches(0.2)
doc.add_paragraph()

# ---------- 본문 렌더 ----------
def render_table(block):
    header = [c.strip() for c in block[0].strip("|").split("|")]
    data = [[c.strip() for c in r.strip("|").split("|")] for r in block[2:]]
    tb = doc.add_table(rows=1, cols=len(header)); tb.style = "Table Grid"
    tb.autofit = True
    for j, htext in enumerate(header):
        c = tb.rows[0].cells[j]; c.text = ""
        pr = c.paragraphs[0]; pr.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        add_runs(pr, htext)
        for run in pr.runs:
            run.bold = True; run.font.size = Pt(10)
    for row in data:
        cells = tb.add_row().cells
        for j, val in enumerate(row):
            if j >= len(cells):
                continue
            cells[j].text = ""
            pr = cells[j].paragraphs[0]; pr.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            add_runs(pr, val)
            for run in pr.runs:
                run.font.size = Pt(10)
    doc.add_paragraph()

def render_figure(s):
    m = re.match(r"\*\*Figure (\d+)\.\*\*\s*(.*)", s)
    num, cap = m.group(1), m.group(2)
    cap = cap.replace("*(caption placeholder)*", "").strip()
    p = doc.add_paragraph()
    p.add_run(f"[Figure {num} 삽입 위치]").bold = True
    if cap:
        p.add_run("  " + cap).italic = True

i0 = i
while i < len(lines):
    raw = lines[i]; s = raw.strip()
    if s.startswith("<!--"):
        while i < len(lines) and "-->" not in lines[i]:
            i += 1
        i += 1; continue
    if s == "" or s == "---":
        i += 1; continue
    if s.startswith("|"):
        block = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            block.append(lines[i].strip()); i += 1
        render_table(block); continue
    if s.startswith("#### "):
        doc.add_heading(clean(s[5:]), level=3); i += 1; continue
    if s.startswith("### "):
        doc.add_heading(clean(s[4:]), level=2); i += 1; continue
    if s.startswith("## "):
        htext = s[3:]
        if "참고문헌" in htext:
            doc.add_page_break()          # 참고문헌 별도 페이지
        doc.add_heading(clean(htext), level=1); i += 1; continue
    if s.startswith("**Figure"):
        render_figure(s); i += 1; continue
    if s.startswith(">"):
        q = []
        while i < len(lines) and lines[i].strip().startswith(">"):
            q.append(lines[i].strip().lstrip(">").strip()); i += 1
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.3)
        add_runs(p, " ".join(x for x in q if x)); continue
    if s.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, s[2:]); i += 1; continue
    p = doc.add_paragraph()
    add_runs(p, s); i += 1; continue

doc.save(OUT)
print("saved:", OUT)
print("본문 시작 라인:", i0, "| 총 처리 라인:", len(lines))
